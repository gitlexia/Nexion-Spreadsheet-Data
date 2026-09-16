from __future__ import annotations

import calendar
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

NAVY = "062D40"
ORANGE = "F58220"
PALE = "F4F7F8"
WHITE = "FFFFFF"
TEXT = "173042"


def write_report(source: str, rows: list[dict], year: int, month: int, output_dir: Path) -> Path:
    month_name = calendar.month_name[month]
    filename = f"nexion_{source}_{year}_{month:02d}_{month_name.lower()}.xlsx"
    path = output_dir / filename

    wb = Workbook()
    ws = wb.active
    ws.title = source.title()

    if source == "facebook":
        columns = [
            "State",
            "Facility",
            "Facebook Likes",
            "Facebook Followers",
            "Facebook Posts (Target Month)",
        ]
    else:
        columns = [
            "State",
            "Facility",
            "Google Rating",
            "Google Review Count",
        ]
        
    # Sort by State A-Z, then Facility A-Z
    rows = sorted(
        rows,
        key=lambda item: (
            (item.get("State") or "").casefold(),
            (item.get("Facility") or "").casefold(),
        ),
    )

    # Title row
    ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=len(columns))
    title = ws.cell(1, 1, f"Nexion {source.title()} Report — {month_name} {year}")
    title.font = Font(size=16, bold=True, color=WHITE)
    title.fill = PatternFill("solid", fgColor=NAVY)
    title.alignment = Alignment(vertical="center")
    ws.row_dimensions[1].height = 30

    # Header row
    for col_idx, name in enumerate(columns, start=1):
        cell = ws.cell(3, col_idx, name)
        cell.font = Font(bold=True, color=WHITE)
        cell.fill = PatternFill("solid", fgColor=ORANGE)
        cell.alignment = Alignment(vertical="center")

    # Data
    for row_idx, item in enumerate(rows, start=4):
        for col_idx, name in enumerate(columns, start=1):
            cell = ws.cell(row_idx, col_idx, item.get(name))
            if row_idx % 2 == 0:
                cell.fill = PatternFill("solid", fgColor=PALE)
            cell.font = Font(color=TEXT)
            cell.alignment = Alignment(vertical="top")

    widths = {
        "State": 10,
        "Facility": 48,
        "Facebook Likes": 18,
        "Facebook Followers": 22,
        "Facebook Posts (Target Month)": 30,
        "Google Rating": 18,
        "Google Review Count": 22,
    }
    for idx, name in enumerate(columns, start=1):
        ws.column_dimensions[get_column_letter(idx)].width = widths.get(name, 18)

    ws.freeze_panes = "A4"
    ws.auto_filter.ref = f"A3:{get_column_letter(len(columns))}{max(3, len(rows) + 3)}"
    ws.sheet_view.showGridLines = False
    wb.save(path)
    return path
